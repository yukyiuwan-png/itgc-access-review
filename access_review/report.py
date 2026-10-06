"""Console summary, Excel workpaper, and a standalone HTML report."""

from __future__ import annotations

import html
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from access_review import __version__
from access_review.engine import Finding, ReviewResult
from access_review.loader import Rule

NAVY = "1B3044"
INK = "1C2430"
BRICK = "9A3412"
AMBER = "92400E"
GREEN = "166534"
HIGH_FILL = "FDE8E4"
MEDIUM_FILL = "FEF3C7"
LOW_FILL = "DCFCE7"
HEADER_FILL = "1B3044"
WHITE = "FFFFFF"

DETAIL_FIELDS = [
    ("rule_id", "Rule ID"),
    ("rule_name", "Test"),
    ("severity", "Severity"),
    ("control_objective", "Control objective"),
    ("iso_controls", "ISO/IEC 27001:2022"),
    ("username", "Username"),
    ("employee_id", "Employee ID"),
    ("employee_name", "Employee name"),
    ("department", "Department"),
    ("manager", "Manager"),
    ("hr_status", "HR status"),
    ("roles", "Roles"),
    ("is_admin", "Admin"),
    ("last_login", "Last login"),
    ("created_date", "Created"),
    ("approval_ticket", "Approval ticket"),
    ("evidence", "Evidence"),
    ("remediation", "Suggested remediation"),
]

COLUMN_WIDTHS = {
    "Rule ID": 14,
    "Test": 46,
    "Severity": 14,
    "Exceptions": 14,
    "Control objective": 46,
    "ISO/IEC 27001:2022": 36,
    "Suggested remediation": 46,
    "Username": 20,
    "Employee ID": 16,
    "Employee name": 22,
    "Department": 24,
    "Manager": 28,
    "HR status": 16,
    "Roles": 36,
    "Admin": 10,
    "Last login": 14,
    "Created": 14,
    "Approval ticket": 18,
    "Evidence": 62,
}

THIN = Border(
    left=Side(style="thin", color="E4D9C8"),
    right=Side(style="thin", color="E4D9C8"),
    top=Side(style="thin", color="E4D9C8"),
    bottom=Side(style="thin", color="E4D9C8"),
)


def format_console(result: ReviewResult, excel_path: Path, html_path: Path) -> str:
    """Plain-text summary an auditor can paste into a workpaper index."""
    objective_width = 72
    lines = [
        "ITGC user access review",
        f"{result.organization} - {result.system_name}",
        f"Review date: {result.review_date.strftime('%Y-%m-%d')}",
        (
            f"Population: {result.employee_count} employees "
            f"({result.active_employees} active, {result.terminated_employees} terminated), "
            f"{result.account_count} accounts"
        ),
        "",
        f"{'Rule':<8} {'Severity':<10} {'Count':>5}  Control objective",
        f"{'----':<8} {'--------':<10} {'-----':>5}  ------------------",
    ]
    for rule in result.enabled_rules:
        count = len(result.for_rule(rule.id))
        objective = rule.control_objective
        if len(objective) > objective_width:
            objective = objective[: objective_width - 3] + "..."
        lines.append(f"{rule.id:<8} {rule.severity:<10} {count:5d}  {objective}")

    lines.append("")
    lines.append(
        f"Exceptions: {len(result.findings)} "
        f"across {result.accounts_with_exceptions} accounts "
        f"(an account can fail more than one test)"
    )
    lines.append(f"Overall rating: {result.overall_rating}")
    if result.disabled_rules:
        skipped = ", ".join(rule.id for rule in result.disabled_rules)
        lines.append(f"Skipped (disabled): {skipped}")
    lines.append(f"Excel: {excel_path}")
    lines.append(f"HTML:  {html_path}")
    return "\n".join(lines)


def _fill_for(severity: str) -> PatternFill:
    colors = {"High": HIGH_FILL, "Medium": MEDIUM_FILL, "Low": LOW_FILL}
    return PatternFill("solid", fgColor=colors.get(severity, WHITE))


def _style_header(cell) -> None:
    cell.font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _apply_widths(sheet: Worksheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS.get(header, 18)


def _sheet_title(rule_id: str, used: set[str]) -> str:
    cleaned = "".join("_" if char in "[]:*?/\\" else char for char in rule_id).strip()
    cleaned = (cleaned or "Rule")[:31]
    candidate = cleaned
    suffix = 2
    while candidate in used:
        tail = f"_{suffix}"
        candidate = cleaned[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate)
    return candidate


def _parameter_note(rule: Rule) -> str:
    if "inactive_days" in rule.parameters:
        return f"Dormant after more than {rule.parameters['inactive_days']} days without a login."
    pairs = rule.parameters.get("conflict_pairs")
    if pairs:
        rendered = ", ".join(f"{left} + {right}" for left, right in pairs)
        return f"Conflict pairs: {rendered}."
    return ""


def _write_summary(sheet: Worksheet, result: ReviewResult) -> None:
    sheet.sheet_properties.tabColor = NAVY
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_view.showGridLines = False

    sheet["A1"] = "ITGC user access review"
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    sheet.merge_cells("A1:G1")

    sheet["A2"] = f"{result.organization}  |  {result.system_name}"
    sheet["A2"].font = Font(name="Calibri", size=12, color=INK)
    sheet.merge_cells("A2:G2")

    meta = [
        ("Review date", result.review_date.strftime("%Y-%m-%d")),
        ("Control domain", result.control_domain),
        (
            "Population",
            (
                f"{result.employee_count} employees "
                f"({result.active_employees} active, {result.terminated_employees} terminated); "
                f"{result.account_count} accounts"
            ),
        ),
        (
            "Exceptions",
            f"{len(result.findings)} exceptions across {result.accounts_with_exceptions} accounts",
        ),
        ("Overall rating", result.overall_rating),
        ("HR source", result.hr_source or "not recorded"),
        ("Account source", result.accounts_source or "not recorded"),
        ("Rules", result.rules_source or "not recorded"),
    ]
    for offset, (label, value) in enumerate(meta):
        row = 4 + offset
        sheet.cell(row, 1, label).font = Font(name="Calibri", bold=True, color="5C6570")
        sheet.cell(row, 2, value).font = Font(name="Calibri", color=INK)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

    header_row = 13
    headers = [
        "Rule ID",
        "Test",
        "Severity",
        "Exceptions",
        "Control objective",
        "ISO/IEC 27001:2022",
        "Suggested remediation",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        _style_header(cell)
    _apply_widths(sheet, headers)
    sheet.row_dimensions[header_row].height = 22
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:G{header_row + max(len(result.enabled_rules), 1)}"

    for index, rule in enumerate(result.enabled_rules):
        row = header_row + 1 + index
        count = len(result.for_rule(rule.id))
        values = [
            rule.id,
            rule.name,
            rule.severity,
            count,
            rule.control_objective,
            rule.iso_text(),
            rule.remediation,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(name="Calibri", size=11, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN
            cell.fill = _fill_for(rule.severity)
        sheet.row_dimensions[row].height = 48

    total_row = header_row + 1 + len(result.enabled_rules)
    sheet.cell(total_row, 1, "Total").font = Font(name="Calibri", bold=True)
    sheet.cell(total_row, 4, len(result.findings)).font = Font(name="Calibri", bold=True)

    note_row = total_row + 2
    note = (
        "Counts are exceptions, not distinct accounts. One account can fail more than one test. "
        "Overall rating is the highest severity among the exceptions. "
        "Satisfactory means this population produced no exceptions. "
        "Severity is set in rules.yaml; it is a review starting point, not a risk assessment. "
        "Every row in the account export is treated as an active account."
    )
    sheet.cell(note_row, 1, note).alignment = Alignment(wrap_text=True)
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    sheet.row_dimensions[note_row].height = 48
    sheet.oddHeader.left.text = "ITGC access review"
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_view.zoomScale = 110


def _write_rule_sheet(sheet: Worksheet, rule: Rule, findings: list[Finding]) -> None:
    count = len(findings)
    if count == 0:
        sheet.sheet_properties.tabColor = GREEN
    elif rule.severity == "High":
        sheet.sheet_properties.tabColor = BRICK
    elif rule.severity == "Medium":
        sheet.sheet_properties.tabColor = AMBER
    else:
        sheet.sheet_properties.tabColor = GREEN

    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4

    last_column = len(DETAIL_FIELDS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet["A1"] = f"{rule.id}  {rule.name}"
    sheet["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)

    details = [
        ("Control objective", rule.control_objective),
        ("Control domain", rule.control_domain),
        ("ISO/IEC 27001:2022", rule.iso_text()),
        ("Severity", rule.severity),
        ("Suggested remediation", rule.remediation),
        ("Exceptions", str(count)),
    ]
    note = _parameter_note(rule)
    if note:
        details.append(("Parameters", note))
    for offset, (label, value) in enumerate(details):
        row = 3 + offset
        sheet.cell(row, 1, label).font = Font(name="Calibri", bold=True, color="5C6570")
        cell = sheet.cell(row, 2, value)
        cell.font = Font(name="Calibri", color=INK)
        cell.alignment = Alignment(wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=min(8, last_column))

    header_row = 3 + len(details) + 1
    headers = [label for _, label in DETAIL_FIELDS]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        _style_header(cell)
    _apply_widths(sheet, headers)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(last_column)}{header_row + max(count, 1)}"
    sheet.print_title_rows = f"{header_row}:{header_row}"

    if count == 0:
        sheet.cell(header_row + 1, 1, "No exceptions. Nothing in this population failed the test.")
        sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=last_column,
        )
        return

    for index, finding in enumerate(findings):
        row = header_row + 1 + index
        for column, (field, _) in enumerate(DETAIL_FIELDS, start=1):
            cell = sheet.cell(row, column, getattr(finding, field))
            cell.font = Font(name="Calibri", size=11, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN
            cell.fill = _fill_for(finding.severity)
        sheet.row_dimensions[row].height = 36


def write_excel(result: ReviewResult, path: Path) -> None:
    """Write one summary sheet and one sheet per enabled rule."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, result)

    used = {"Summary"}
    for rule in result.enabled_rules:
        sheet = workbook.create_sheet(_sheet_title(rule.id, used))
        _write_rule_sheet(sheet, rule, result.for_rule(rule.id))

    workbook.properties.title = "ITGC user access review"
    workbook.properties.creator = f"itgc-access-review {__version__}"
    workbook.properties.subject = result.control_domain
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _esc(value: object) -> str:
    return html.escape(str(value), quote=True)


def _blank(value: str) -> str:
    return _esc(value) if value else "—"


def _fact(label: str, value: str) -> str:
    return (
        "<div>"
        f"<dt>{_esc(label)}</dt>"
        f"<dd>{_blank(value)}</dd>"
        "</div>"
    )


def _finding_card(finding: Finding) -> str:
    who_bits = [
        finding.employee_name or "No HR record",
        finding.employee_id or "(blank id)",
    ]
    if finding.department:
        who_bits.append(finding.department)
    if finding.manager:
        who_bits.append(f"Manager {finding.manager}")
    who = " · ".join(who_bits)
    severity = finding.severity.lower()
    facts = "".join(
        [
            _fact("HR status", finding.hr_status),
            _fact("Roles", finding.roles),
            _fact("Admin", finding.is_admin),
            _fact("Last login", finding.last_login),
            _fact("Created", finding.created_date),
            _fact("Approval ticket", finding.approval_ticket),
        ]
    )
    return f"""
    <article class="finding {severity}">
      <header>
        <code>{_esc(finding.username)}</code>
        <span class="pill {severity}">{_esc(finding.severity)}</span>
      </header>
      <p class="who">{_esc(who)}</p>
      <p class="evidence">{_esc(finding.evidence)}</p>
      <dl class="facts">{facts}</dl>
      <p class="fix"><span>Remediation</span> {_esc(finding.remediation)}</p>
    </article>
    """


def _summary_rows(result: ReviewResult) -> str:
    rows = []
    for rule in result.enabled_rules:
        count = len(result.for_rule(rule.id))
        rows.append(
            "<tr>"
            f"<td><a href=\"#{_esc(rule.id)}\"><code>{_esc(rule.id)}</code></a></td>"
            f"<td>{_esc(rule.name)}</td>"
            f"<td><span class=\"pill {rule.severity.lower()}\">{_esc(rule.severity)}</span></td>"
            f"<td class=\"num\">{count}</td>"
            f"<td>{_esc(rule.control_objective)}</td>"
            f"<td>{_esc(rule.iso_text())}</td>"
            "</tr>"
        )
    rows.append(
        "<tr class=\"total\">"
        "<td></td><td>Total exceptions</td><td></td>"
        f"<td class=\"num\">{len(result.findings)}</td><td></td><td></td>"
        "</tr>"
    )
    return "\n".join(rows)


def _rule_section(result: ReviewResult, rule: Rule) -> str:
    findings = result.for_rule(rule.id)
    note = _parameter_note(rule)
    note_html = f"<p class=\"meta\">{_esc(note)}</p>" if note else ""
    if findings:
        body = "\n".join(_finding_card(finding) for finding in findings)
    else:
        body = '<p class="empty">No exceptions. Nothing in this population failed the test.</p>'
    return f"""
    <section id="{_esc(rule.id)}">
      <h2><code>{_esc(rule.id)}</code> {_esc(rule.name)}</h2>
      <p class="objective">{_esc(rule.control_objective)}</p>
      <p class="meta">{_esc(rule.control_domain)} · ISO/IEC 27001:2022 {_esc(rule.iso_text())}</p>
      {note_html}
      <p class="fix section-fix"><span>Suggested remediation</span> {_esc(rule.remediation)}</p>
      <p class="meta">{len(findings)} exception{"s" if len(findings) != 1 else ""}</p>
      {body}
    </section>
    """


def render_html(result: ReviewResult, generated_at: datetime | None = None) -> str:
    """Return a self-contained HTML workpaper."""
    stamp = (generated_at or datetime.now()).strftime("%Y-%m-%d %H:%M")
    as_of = result.review_date.strftime("%d %B %Y")
    sources = ""
    if result.hr_source or result.accounts_source or result.rules_source:
        sources = (
            "<p class=\"meta\">Sources: "
            f"HR {_esc(result.hr_source or 'n/a')}, "
            f"accounts {_esc(result.accounts_source or 'n/a')}, "
            f"rules {_esc(result.rules_source or 'n/a')}.</p>"
        )
    rating = result.overall_rating
    rating_class = "ok" if rating == "Satisfactory" else rating.lower()
    sections = "\n".join(_rule_section(result, rule) for rule in result.enabled_rules)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>ITGC access review — {_esc(result.organization)}</title>
  <style>
    :root {{
      --ink: #1c2430;
      --muted: #5c6570;
      --paper: #f3efe6;
      --card: #fffcf7;
      --line: #e2d9cc;
      --navy: #1b3044;
      --brick: #9a3412;
      --amber: #92400e;
      --green: #166534;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--paper);
      color: var(--ink);
      font: 15px/1.5 "Segoe UI", "Helvetica Neue", Arial, sans-serif;
    }}
    .banner {{ background: var(--navy); color: #f7f4ee; }}
    .banner-inner, main {{ max-width: 1040px; margin: 0 auto; padding: 32px 20px; }}
    main {{ padding-top: 28px; padding-bottom: 64px; }}
    .kicker {{
      margin: 0 0 8px;
      letter-spacing: 0.14em;
      text-transform: uppercase;
      font-size: 12px;
      color: #d7c4a3;
    }}
    h1, h2 {{
      font-family: Georgia, "Iowan Old Style", Palatino, "Palatino Linotype", serif;
      font-weight: normal;
    }}
    h1 {{ margin: 0; font-size: 36px; line-height: 1.15; }}
    h2 {{ margin: 36px 0 8px; font-size: 26px; }}
    h2 code {{ font-size: 0.75em; }}
    .sub {{ margin: 10px 0 0; color: #e7e1d6; }}
    .stats {{
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      gap: 12px;
      margin-top: -28px;
    }}
    .stat {{
      background: var(--card);
      border: 1px solid var(--line);
      padding: 12px 14px;
    }}
    .stat span {{ display: block; color: var(--muted); font-size: 12px; letter-spacing: 0.04em; text-transform: uppercase; }}
    .stat b {{ display: block; margin-top: 2px; font-family: Georgia, Palatino, serif; font-size: 28px; font-weight: normal; }}
    .stat.high b {{ color: var(--brick); }}
    .stat.medium b {{ color: var(--amber); }}
    .stat.ok b {{ color: var(--green); }}
    .note {{ color: var(--muted); font-size: 14px; }}
    .table-scroll {{ overflow-x: auto; border: 1px solid var(--line); background: var(--card); }}
    table {{ width: 100%; border-collapse: collapse; min-width: 720px; }}
    th {{
      text-align: left;
      font-size: 12px;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      color: var(--muted);
      border-bottom: 2px solid var(--navy);
      padding: 10px;
    }}
    td {{ border-bottom: 1px solid var(--line); vertical-align: top; padding: 10px; }}
    td.num, th.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
    tr.total td {{ font-weight: 600; border-bottom: none; }}
    a {{ color: var(--navy); }}
    code {{ font-family: ui-monospace, "SFMono-Regular", Menlo, Consolas, monospace; font-size: 0.92em; }}
    .pill {{
      display: inline-block;
      padding: 1px 8px;
      border-radius: 999px;
      font-size: 12px;
      letter-spacing: 0.03em;
      text-transform: uppercase;
    }}
    .pill.high {{ background: #fde8e4; color: var(--brick); }}
    .pill.medium {{ background: #fef3c7; color: var(--amber); }}
    .pill.low, .pill.ok {{ background: #dcfce7; color: var(--green); }}
    .objective {{ margin: 0 0 4px; }}
    .meta {{ margin: 0 0 8px; color: var(--muted); font-size: 14px; }}
    .finding {{
      background: var(--card);
      border: 1px solid var(--line);
      border-left: 4px solid var(--navy);
      padding: 12px 14px;
      margin: 10px 0;
    }}
    .finding.high {{ border-left-color: var(--brick); }}
    .finding.medium {{ border-left-color: var(--amber); }}
    .finding.low {{ border-left-color: var(--green); }}
    .finding header {{ display: flex; justify-content: space-between; gap: 12px; align-items: center; }}
    .who {{ margin: 4px 0; }}
    .evidence {{ margin: 6px 0; }}
    .facts {{ display: flex; flex-wrap: wrap; gap: 8px 18px; margin: 8px 0 0; }}
    .facts div {{ min-width: 8.5rem; }}
    .facts dt {{
      font-size: 11px;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      color: var(--muted);
    }}
    .facts dd {{ margin: 0; }}
    .fix {{ margin: 8px 0 0; font-size: 14px; }}
    .fix span, .section-fix span {{
      display: inline-block;
      margin-right: 6px;
      font-size: 11px;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      color: var(--muted);
    }}
    .empty {{
      background: var(--card);
      border: 1px dashed var(--line);
      padding: 12px 14px;
      color: var(--green);
    }}
    footer {{
      max-width: 1040px;
      margin: 0 auto;
      padding: 0 20px 40px;
      color: var(--muted);
      font-size: 13px;
    }}
    @media (max-width: 720px) {{
      h1 {{ font-size: 28px; }}
      .stats {{ grid-template-columns: 1fr 1fr; }}
    }}
    @media print {{
      body {{ background: white; }}
      .banner {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
      .finding, .stat {{ break-inside: avoid; }}
      a {{ text-decoration: none; color: inherit; }}
    }}
  </style>
</head>
<body>
  <header class="banner">
    <div class="banner-inner">
      <p class="kicker">IT general controls · {_esc(result.control_domain)}</p>
      <h1>User access review</h1>
      <p class="sub">{_esc(result.organization)} · {_esc(result.system_name)} · As of {_esc(as_of)}</p>
    </div>
  </header>
  <main>
    <section class="stats" aria-label="Review totals">
      <div class="stat"><span>Employees</span><b>{result.employee_count}</b></div>
      <div class="stat"><span>Accounts</span><b>{result.account_count}</b></div>
      <div class="stat"><span>Exceptions</span><b>{len(result.findings)}</b></div>
      <div class="stat {rating_class}"><span>Overall rating</span><b>{_esc(rating)}</b></div>
    </section>
    <p class="note">
      {result.active_employees} active and {result.terminated_employees} terminated employees.
      {result.accounts_with_exceptions} accounts have at least one exception.
      An account can appear under more than one test. Severity comes from the rule file;
      it is a starting point for the reviewer, not a risk rating.
    </p>
    {sources}
    <h2>Summary</h2>
    <div class="table-scroll">
      <table>
        <thead>
          <tr>
            <th>Rule</th>
            <th>Test</th>
            <th>Severity</th>
            <th class="num">Exceptions</th>
            <th>Control objective</th>
            <th>ISO/IEC 27001:2022</th>
          </tr>
        </thead>
        <tbody>
          {_summary_rows(result)}
        </tbody>
      </table>
    </div>
    {sections}
  </main>
  <footer>
    <p>
      Generated by itgc-access-review {__version__} on {stamp}.
      Sample populations are fictional and are not an engagement deliverable.
      The tool checks ticket presence, not ticket validity, and it treats every
      exported account as active.
    </p>
  </footer>
</body>
</html>
"""


def write_html(result: ReviewResult, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(result), encoding="utf-8")
