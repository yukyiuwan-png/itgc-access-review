"""Excel, HTML, and the command line, including a test with nothing to report."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from access_review.cli import main
from access_review.engine import run_review
from access_review.loader import Config, Rule
from access_review.report import render_html, write_excel

ROOT = Path(__file__).resolve().parents[1]


def _rule(**overrides) -> Rule:
    data = {
        "id": "TER-01",
        "type": "terminated_with_access",
        "name": "Terminated employee still has an active account",
        "enabled": True,
        "severity": "High",
        "control_objective": "Access is removed when employment ends.",
        "control_domain": "Access to programs and data",
        "iso_controls": ["5.18 Access rights", "5.15 Access control"],
        "remediation": "Disable the account.",
        "parameters": {},
    }
    data.update(overrides)
    return Rule(**data)


def _clean_config() -> Config:
    return Config(
        organization="Sample Org (fictional)",
        system_name="LedgerOne",
        review_date=pd.Timestamp("2026-09-30"),
        control_domain="Access to programs and data",
        rules=[_rule()],
    )


def test_html_states_when_a_rule_finds_nothing():
    result = run_review(
        pd.DataFrame(
            [
                {
                    "employee_id": "E1",
                    "name": "Mina Park",
                    "department": "Human Resources",
                    "status": "active",
                    "termination_date": "",
                    "manager": "Alex Rowan (Director)",
                }
            ]
        ),
        pd.DataFrame(
            [
                {
                    "username": "mpark",
                    "employee_id": "E1",
                    "roles": "view_reports",
                    "is_admin": "FALSE",
                    "last_login": "2026-09-20",
                    "created_date": "2026-09-01",
                    "approval_ticket": "",
                }
            ]
        ),
        _clean_config(),
    )
    page = render_html(result, generated_at=pd.Timestamp("2026-10-06 09:00").to_pydatetime())
    assert "No exceptions" in page
    assert "Satisfactory" in page
    assert "TER-01" in page
    assert "5.18 Access rights" in page


def test_excel_has_a_summary_and_one_sheet_per_rule(tmp_path):
    from access_review.loader import load_accounts, load_hr, load_rules

    result = run_review(
        load_hr(ROOT / "sample_data" / "hr.csv"),
        load_accounts(ROOT / "sample_data" / "accounts.csv"),
        load_rules(ROOT / "rules.yaml"),
        hr_source="sample_data/hr.csv",
        accounts_source="sample_data/accounts.csv",
        rules_source="rules.yaml",
    )
    path = tmp_path / "review.xlsx"
    write_excel(result, path)
    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Summary"
    assert workbook.sheetnames[1:] == ["TER-01", "ORP-01", "ADM-01", "DRM-01", "SOD-01", "PST-01"]

    summary_text = " ".join(
        str(cell.value)
        for row in workbook["Summary"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "TER-01" in summary_text
    assert "High" in summary_text
    assert "29" in summary_text

    ter_text = " ".join(
        str(cell.value)
        for row in workbook["TER-01"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "mhale" in ter_text
    assert "Disable or delete the account" in ter_text


def test_cli_writes_both_reports(tmp_path, capsys):
    code = main(
        [
            "--hr",
            str(ROOT / "sample_data" / "hr.csv"),
            "--accounts",
            str(ROOT / "sample_data" / "accounts.csv"),
            "--rules",
            str(ROOT / "rules.yaml"),
            "--out",
            str(tmp_path),
        ]
    )
    assert code == 0
    assert (tmp_path / "itgc_access_review.xlsx").is_file()
    assert (tmp_path / "itgc_access_review.html").is_file()
    output = capsys.readouterr().out
    assert "Overall rating: High" in output
    assert "TER-01" in output
    assert "DRM-01" in output


def test_cli_returns_an_error_when_a_file_is_missing(tmp_path, capsys):
    code = main(
        [
            "--hr",
            str(tmp_path / "missing.csv"),
            "--accounts",
            str(ROOT / "sample_data" / "accounts.csv"),
            "--rules",
            str(ROOT / "rules.yaml"),
            "--out",
            str(tmp_path),
        ]
    )
    assert code == 1
    assert "File not found" in capsys.readouterr().err
